"""Editable Word (.docx) versions of the student pages (+ lesson plan for Pro).

Emoji and diagrams are inserted as PNG pictures so they look the same in
Word, Google Docs and LibreOffice (colour-emoji text support varies).
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import config, illustrations, mathgen, spec as S


def _rgb(hexcode):
    h = hexcode.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell, hexcode):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcode.lstrip("#"))
    tc.append(shd)


class Doc:
    def __init__(self, spec):
        self.spec = spec
        self.color = config.SUBJECT_COLORS[spec["meta"]["subject"]]
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(1.6))
        st = self.d.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(12)
        for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
            h = self.d.styles[name]
            h.font.name = "Arial"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = _rgb(config.BRAND["navy"])
        footer = sec.footer.paragraphs[0]
        footer.text = f"TeachLuma · {spec['meta']['title']} · teachluma.com"
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = _rgb("#8A93A8")
        self.first = True

    def page(self, kicker, title):
        if not self.first:
            self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.first = False
        p = self.d.add_paragraph()
        r = p.add_run(kicker.upper())
        r.font.size, r.font.bold, r.font.color.rgb = Pt(9), True, _rgb(self.color["main"])
        self.d.add_heading(title, level=1)

    def namebar(self, score=None):
        p = self.d.add_paragraph("Name: ______________________________   Date: ____________" + (f"   Score: ____ / {score}" if score else ""))
        p.runs[0].font.bold = True

    def para(self, text, bold=False, italic=False, color=None, size=None):
        p = self.d.add_paragraph()
        r = p.add_run(str(text))
        r.bold, r.italic = bold, italic
        if color:
            r.font.color.rgb = _rgb(color)
        if size:
            r.font.size = Pt(size)
        return p

    def bullets(self, lines, style="List Bullet"):
        for x in S.as_lines(lines):
            self.d.add_paragraph(x, style=style)

    def picture(self, run, path, cm):
        run.add_picture(path, width=Cm(cm))

    def cell_content(self, cell, value, img_cm=1.3, prefix=""):
        cell.text = ""
        p = cell.paragraphs[0]
        if prefix:
            p.add_run(prefix + "  ")
        if illustrations.is_emoji(str(value)):
            self.picture(p.add_run(), illustrations.emoji_png(str(value), 160), img_cm)
        else:
            p.add_run(str(value)).bold = True

    def table(self, rows, header=None, widths=None):
        t = self.d.add_table(rows=0, cols=len(rows[0] if rows else header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if header:
            cells = t.add_row().cells
            for c, h in zip(cells, header):
                c.text = str(h)
                c.paragraphs[0].runs[0].bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                _shade(c, config.BRAND["navy"])
        for row in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                self.cell_content(c, v) if illustrations.is_emoji(str(v)) else setattr(c, "text", str(v))
        if widths:
            self._widths(t, widths)
        self.d.add_paragraph()
        return t

    def _widths(self, t, widths):
        # Word reads cell widths, LibreOffice reads the grid columns: set both.
        t.autofit = False
        for col, w in zip(t.columns, widths):
            col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Cm(w)

    # ------------------------------------------------------------ sections
    def section(self, sec, n, seed):
        t = sec["type"]
        title = sec.get("title") or t.replace("_", " ").capitalize()
        self.d.add_heading(f"{n}. {title}", level=2)
        self.para(sec["instructions"], italic=True, color=config.BRAND["muted"])
        if sec.get("passage"):
            if sec.get("passage_title"):
                self.para(sec["passage_title"], bold=True, color=self.color["ink"], size=13)
            for line in S.as_lines(sec["passage"]):
                self.para(line, size=12)
        bank = S.word_bank(sec, seed)
        if bank:
            self.para(("Words: " if t == "sort" else "Word bank: ") + "   ·   ".join(bank), bold=True, color=self.color["ink"])
        items = sec.get("items", [])
        if t == "match":
            rights = S.shuffled([it["right"] for it in items], seed)
            t_ = self.d.add_table(rows=0, cols=3)
            for i, it in enumerate(items):
                cells = t_.add_row().cells
                self.cell_content(cells[0], it["left"], prefix=f"{i + 1}.")
                cells[1].text = "•          •"
                self.cell_content(cells[2], rights[i], prefix=f"{S.LETTERS[i]}.")
            self._widths(t_, [6, 3, 8])
            self.d.add_paragraph()
        elif t == "fill_blank":
            for i, it in enumerate(items):
                self.d.add_paragraph(f"{i + 1}. " + it["text"].replace("____", "__________"))
        elif t == "multiple_choice":
            for i, it in enumerate(items):
                p = self.d.add_paragraph(f"{i + 1}. ")
                if illustrations.is_emoji(str(it["q"])):
                    self.picture(p.add_run(), illustrations.emoji_png(it["q"], 160), 1.2)
                else:
                    p.add_run(str(it["q"]))
                p = self.d.add_paragraph("      ")
                for k, o in enumerate(it["options"]):
                    p.add_run(f"○ {S.LETTERS[k]}) ")
                    if illustrations.is_emoji(str(o)):
                        self.picture(p.add_run(), illustrations.emoji_png(str(o), 160), 1.0 * max(1, len(str(o)) if not str(o).startswith("pic:") else 1))
                    else:
                        p.add_run(str(o))
                    p.add_run("        ")
        elif t == "true_false":
            for i, it in enumerate(items):
                self.d.add_paragraph(f"{i + 1}. {it['statement']}      ○ True   ○ False")
        elif t == "short_answer":
            for i, it in enumerate(items):
                self.d.add_paragraph(f"{i + 1}. {it['q']}")
                for _ in range(sec.get("lines", 2)):
                    self.d.add_paragraph("_" * 70)
        elif t == "label":
            cols = 4
            t_ = self.d.add_table(rows=0, cols=cols)
            t_.style = "Table Grid"
            for i in range(0, len(items), cols):
                cells = t_.add_row().cells
                for j, it in enumerate(items[i:i + cols]):
                    cells[j].text = f"{i + j + 1}"
                    self.picture(cells[j].add_paragraph().add_run(), illustrations.emoji_png(it["emoji"], 160), 1.8)
                    cells[j].add_paragraph("____________")
            self.d.add_paragraph()
        elif t == "sequence":
            order = S.shuffled(items, seed)
            self.table([["", f"{S.LETTERS[i]}. {s}"] for i, s in enumerate(order)], widths=[1.5, 15.5])
        elif t == "sort":
            groups = list(sec["groups"])
            rows = max(len(v) for v in sec["groups"].values())
            self.table([[""] * len(groups) for _ in range(rows)], header=groups)
        elif t == "draw_write":
            for i, it in enumerate(items):
                self.d.add_paragraph(f"{i + 1}. {it['prompt']}").paragraph_format.keep_with_next = True
                if it.get("box", True):
                    box = self.d.add_table(rows=1, cols=1)
                    box.style = "Table Grid"
                    box.rows[0].height = Cm(5.5)
                for _ in range(max(1, it.get("lines", 1))):
                    self.d.add_paragraph("_" * 70)
        elif t == "math":
            graph = next(((it.get("render") or {}).get("pictograph") for it in items if (it.get("render") or {}).get("pictograph")), None)
            if graph:
                title, cats = mathgen.PICTO_THEMES[graph[0]]
                self.para(title, bold=True, color=config.BRAND["navy"])
                g = self.d.add_table(rows=0, cols=2)
                g.style = "Table Grid"
                for (icon, name), n in zip(cats, mathgen.picto_data(*graph)):
                    cells = g.add_row().cells
                    cells[0].text = name.capitalize()
                    run = cells[1].paragraphs[0].add_run()
                    for _ in range(n):
                        self.picture(run, illustrations.emoji_png(icon, 120), 0.7)
                self._widths(g, [3.5, 13.5])
                self.para("Each picture shows 1.", italic=True, color=config.BRAND["muted"], size=9)
            plot = next(((it.get("render") or {}).get("lineplot") for it in items if (it.get("render") or {}).get("lineplot")), None)
            if plot:
                title, things, unit = mathgen.LINEPLOT_THEMES[plot[0]]
                self.para(title, bold=True, color=config.BRAND["navy"])
                self.picture(self.d.add_paragraph().add_run(), illustrations.lineplot_png(mathgen.lineplot_data(*plot), f"{unit}s"), 15)
                self.para(f"Each X shows one of the {things}.", italic=True, color=config.BRAND["muted"], size=9)
            cols = 2
            t_ = self.d.add_table(rows=0, cols=cols)
            for i in range(0, len(items), cols):
                cells = t_.add_row().cells
                for j, it in enumerate(items[i:i + cols]):
                    cells[j].text = f"{i + j + 1}.  " + it["prompt"].replace("____", "________")
                    r = it.get("render") or {}
                    if "clock" in r:
                        self.picture(cells[j].add_paragraph().add_run(), illustrations.clock_png(r["clock"][0], r["clock"][1], 300, r.get("hands", True)), 3.4)
                    elif "fraction" in r:
                        self.picture(cells[j].add_paragraph().add_run(), illustrations.fraction_png(*r["fraction"], size=300), 3.2)
                    elif "thermo" in r:
                        self.picture(cells[j].add_paragraph().add_run(), illustrations.thermometer_png(r["thermo"]), 1.8)
                    elif "objects" in r:
                        run = cells[j].add_paragraph().add_run()
                        for _ in range(r["objects"][1]):
                            self.picture(run, illustrations.emoji_png(r["objects"][0], 120), 0.75)
                    elif "base10" in r:
                        png = illustrations.base10_png(*r["base10"])
                        from PIL import Image
                        with Image.open(png) as im:
                            self.picture(cells[j].add_paragraph().add_run(), png, min(7.5, im.width / 40))
                    cells[j].add_paragraph()
            self.d.add_paragraph()

    def part(self, part):
        m, doc = self.spec["meta"], self.spec[part]
        self.page(f"{S.PART_TITLES[part]} · {m['grade']}", doc.get("title") or m["topic"])
        self.namebar(_points(doc) if part == "quiz" else None)
        for i, sec in enumerate(doc["sections"]):
            self.section(sec, i + 1, S.seed_for(self.spec, part, i))

    def lesson_plan(self):
        m, lp = self.spec["meta"], self.spec["lesson_plan"]
        self.page("Teacher lesson plan", m["title"])
        self.para(f"{m['grade']} · {config.SUBJECT_COLORS[m['subject']]['name']} · {m['suggested_duration']}", bold=True, color=self.color["ink"])
        self.para(m["overview"])
        self.d.add_heading("Learning objectives", level=2)
        self.bullets(m["learning_objectives"])
        self.d.add_heading("Materials", level=2)
        self.bullets(lp["materials"])
        if self.spec.get("vocabulary"):
            self.d.add_heading("Key vocabulary", level=2)
            self.table([[v.get("emoji", ""), v["word"], v["definition"], v.get("example", "")] for v in self.spec["vocabulary"]],
                       header=["", "Word", "Meaning", "Example"], widths=[1.6, 3.9, 5.8, 5.7])
        self.d.add_heading("Lesson procedure", level=2)
        self.table([[st["name"], f"{st['minutes']} min", "\n".join("• " + x for x in S.as_lines(st["steps"]))] for st in lp["stages"]],
                   header=["Stage", "Time", "What to do"], widths=[3.2, 2, 11.8])
        diff = lp.get("differentiation") or {}
        for k, lbl in (("support", "Support"), ("challenge", "Challenge")):
            if diff.get(k):
                self.d.add_heading(lbl, level=2)
                self.bullets(diff[k])
        if lp.get("assessment"):
            self.d.add_heading("Assessment", level=2)
            self.para(lp["assessment"])

    def rubric(self):
        r = self.spec["rubric"]
        self.page("Assessment rubric", self.spec["meta"]["topic"])
        self.table([[c["name"], *c["levels"]] for c in r["criteria"]], header=["Criteria", "4 · Excellent", "3 · Good", "2 · Developing", "1 · Beginning"])

    def answer_key(self):
        self.page("Answer key", self.spec["meta"]["title"])
        for part in S.parts(self.spec):
            self.d.add_heading(S.PART_TITLES[part], level=2)
            for i, sec in enumerate(self.spec[part]["sections"]):
                self.para(f"{i + 1}. {sec.get('title') or sec['type'].replace('_', ' ').capitalize()}", bold=True, color=self.color["ink"])
                script = S.as_lines(sec.get("teacher_script"))
                if script:
                    self.para("Teacher reads aloud:", bold=True)
                    for line in script:
                        self.d.add_paragraph(line, style="List Bullet")
                for a, b in S.section_answers(sec, S.seed_for(self.spec, part, i)):
                    self.d.add_paragraph(f"{a}   {b}")

    def save(self, out):
        self.d.core_properties.title = self.spec["meta"]["title"]
        self.d.core_properties.author = "TeachLuma"
        # The library's template ships its own comment and last-modified name; clear them.
        self.d.core_properties.last_modified_by = "TeachLuma"
        self.d.core_properties.comments = ""
        self.d.core_properties.keywords = ""
        self.d.core_properties.subject = ""
        self.d.core_properties.category = ""
        self.d.core_properties.revision = 1
        self.d.core_properties.created = self.d.core_properties.modified = datetime.now()
        self.d.save(out)
        return Path(out)


def _points(doc):
    return sum(len(s.get("items", [])) or sum(len(v) for v in s.get("groups", {}).values()) for s in doc["sections"])


def build(spec, out):
    d = Doc(spec)
    if spec["meta"]["tier"] == "pro":
        d.lesson_plan()
    for part in S.parts(spec):
        d.part(part)
    if spec["meta"]["tier"] == "pro":
        d.rubric()
    d.answer_key()
    return d.save(out)
